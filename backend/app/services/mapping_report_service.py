"""
CO-PO-PSO Mapping Report Service

This service is used ONLY by Agent 1:
    CO-PO-PSO Mapping Agent

It must not calculate or display attainment.
Agent 2 continues using services/report_service.py.
"""

import os
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle


THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _clean(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _mapping_value(value: Any) -> str:
    """Render mapping values as 1/2/3 or blank."""
    if value in (None, "", 0, "0"):
        return ""

    try:
        number = int(value)
        if number in (1, 2, 3):
            return str(number)
    except (TypeError, ValueError):
        pass

    text = _clean(value).lower()
    if text in ("low",):
        return "1"
    if text in ("moderate", "medium"):
        return "2"
    if text in ("high",):
        return "3"

    return ""


def _normalize_rows(
    rows: Any,
    co_codes: List[str],
    outcome_codes: List[str],
) -> Dict[str, Dict[str, str]]:
    """
    Accept the mapping-agent structure:

    [
        {
            "co_code": "CO1",
            "mappings": {"PO1": 3, "PO2": 2}
        }
    ]

    and return a rectangular lookup.
    """
    lookup = {
        str(co).upper(): {str(code).upper(): "" for code in outcome_codes}
        for co in co_codes
    }

    if not isinstance(rows, list):
        return lookup

    for row in rows:
        if not isinstance(row, dict):
            continue

        co = row.get("co_code") or row.get("co") or row.get("CO")
        if not co:
            continue

        co = str(co).strip().upper()
        if co not in lookup:
            continue

        mappings = row.get("mappings", {})
        if not isinstance(mappings, dict):
            continue

        for outcome in outcome_codes:
            if outcome in mappings:
                lookup[co][outcome] = _mapping_value(mappings[outcome])
            else:
                # case-insensitive lookup
                for key, value in mappings.items():
                    if str(key).strip().upper() == str(outcome).upper():
                        lookup[co][outcome] = _mapping_value(value)
                        break

    return lookup


def _extract_codes(result: Dict[str, Any], key: str) -> List[str]:
    values = result.get(key, [])
    if not isinstance(values, list):
        return []

    output = []
    for value in values:
        code = _clean(value).upper()
        if code and code not in output:
            output.append(code)
    return output


def _table_rows(
    lookup: Dict[str, Dict[str, str]],
    co_codes: List[str],
    outcome_codes: List[str],
) -> List[List[str]]:
    rows = [[
        "CO",
        *outcome_codes,
    ]]

    for co in co_codes:
        rows.append([
            co,
            *[
                lookup.get(co, {}).get(outcome, "")
                for outcome in outcome_codes
            ],
        ])

    return rows


def _style_excel_table(ws, start_row: int, end_row: int, end_col: int):
    for row in ws.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=1,
        max_col=end_col,
    ):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    for cell in ws[start_row]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )


def _write_excel_mapping_sheet(
    ws,
    title: str,
    lookup: Dict[str, Dict[str, str]],
    co_codes: List[str],
    outcome_codes: List[str],
):
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max(1, len(outcome_codes) + 1),
    )

    ws.append([])
    rows = _table_rows(
        lookup,
        co_codes,
        outcome_codes,
    )

    for row in rows:
        ws.append(row)

    end_row = ws.max_row
    end_col = len(outcome_codes) + 1

    _style_excel_table(
        ws,
        3,
        end_row,
        end_col,
    )

    for col in range(1, end_col + 1):
        ws.column_dimensions[
            get_column_letter(col)
        ].width = 16

    ws.column_dimensions["A"].width = 14
    ws.freeze_panes = "B4"


def _mapping_data(result: Dict[str, Any]):
    co_codes = _extract_codes(result, "co_codes")
    po_codes = _extract_codes(result, "po_codes")
    pso_codes = _extract_codes(result, "pso_codes")

    co_po = _normalize_rows(
        result.get("co_po_mapping", []),
        co_codes,
        po_codes,
    )

    co_pso = _normalize_rows(
        result.get("co_pso_mapping", []),
        co_codes,
        pso_codes,
    )

    return (
        co_codes,
        po_codes,
        pso_codes,
        co_po,
        co_pso,
    )


def generate_mapping_excel(
    result: Dict[str, Any],
    output_path: str,
):
    """
    Generate ONLY the Agent 1 mapping Excel report.
    """
    (
        co_codes,
        po_codes,
        pso_codes,
        co_po,
        co_pso,
    ) = _mapping_data(result)

    wb = Workbook()
    ws = wb.active
    ws.title = "CO-PO Mapping"

    _write_excel_mapping_sheet(
        ws,
        "CO-PO Mapping",
        co_po,
        co_codes,
        po_codes,
    )

    ws2 = wb.create_sheet("CO-PSO Mapping")

    _write_excel_mapping_sheet(
        ws2,
        "CO-PSO Mapping",
        co_pso,
        co_codes,
        pso_codes,
    )

    ws3 = wb.create_sheet("Mapping Scale")
    ws3.append(["Mapping Level", "Meaning"])
    ws3.append([1, "Low"])
    ws3.append([2, "Moderate"])
    ws3.append([3, "High"])
    ws3.append(["Blank", "No meaningful contribution"])

    for row in ws3.iter_rows(
        min_row=1,
        max_row=5,
        min_col=1,
        max_col=2,
    ):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    for cell in ws3[1]:
        cell.font = Font(bold=True)

    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 32

    os.makedirs(
        os.path.dirname(output_path) or ".",
        exist_ok=True,
    )

    wb.save(output_path)
    return output_path


def _pdf_table(
    rows: List[List[Any]],
    col_widths=None,
):
    table = Table(
        rows,
        colWidths=col_widths,
        repeatRows=1,
        hAlign="LEFT",
    )

    table.setStyle(
        TableStyle([
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.6,
                colors.black,
            ),
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#D9EAF7"),
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER",
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE",
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                8.5,
            ),
            (
                "LEFTPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
            (
                "RIGHTPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                5,
            ),
        ])
    )

    return table


def generate_mapping_pdf(
    result: Dict[str, Any],
    output_path: str,
):
    """
    Generate ONLY the Agent 1 CO-PO-PSO Mapping PDF.

    No attainment information is included.
    """
    (
        co_codes,
        po_codes,
        pso_codes,
        co_po,
        co_pso,
    ) = _mapping_data(result)

    course_code = _clean(
        result.get("course_code")
    )
    course_name = _clean(
        result.get("course_name")
    )
    mapping_id = _clean(
        result.get("mapping_id")
    )
    approval_status = _clean(
        result.get("approval_status")
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]
    title_style.fontName = "Helvetica-Bold"
    title_style.fontSize = 17
    title_style.leading = 21
    title_style.alignment = 1

    heading_style = styles["Heading2"]
    heading_style.fontName = "Helvetica-Bold"
    heading_style.fontSize = 12
    heading_style.leading = 15

    normal_style = styles["BodyText"]
    normal_style.fontSize = 9.5
    normal_style.leading = 13

    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title="CO-PO-PSO Mapping Report",
        author="FACULTY AGENT",
    )

    story = []

    story.append(
        Table(
            [["FACULTY AGENT"]],
            colWidths=[260 * mm],
            style=TableStyle([
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    16,
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    8,
                ),
            ]),
        )
    )

    story.append(
        Table(
            [["CO-PO-PSO Mapping Report"]],
            colWidths=[260 * mm],
            style=TableStyle([
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    14,
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    10,
                ),
            ]),
        )
    )

    course_info = [
        ["Course Code", course_code],
        ["Course Name", course_name],
        ["Mapping ID", mapping_id],
        ["Approval Status", approval_status],
    ]

    info_table = Table(
        course_info,
        colWidths=[45 * mm, 95 * mm],
        hAlign="LEFT",
    )

    info_table.setStyle(
        TableStyle([
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.6,
                colors.black,
            ),
            (
                "FONTNAME",
                (0, 0),
                (0, -1),
                "Helvetica-Bold",
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "LEFT",
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE",
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                9,
            ),
        ])
    )

    story.append(info_table)
    story.append(Spacer(1, 7 * mm))

    # --------------------------------------------------------
    # CO-PO
    # --------------------------------------------------------

    story.append(
        Table(
            [["CO-PO Mapping"]],
            colWidths=[260 * mm],
            style=TableStyle([
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    12,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]),
        )
    )

    po_rows = _table_rows(
        co_po,
        co_codes,
        po_codes,
    )

    if len(po_rows[0]) > 1:
        width = 255 * mm / len(po_rows[0])
        story.append(
            _pdf_table(
                po_rows,
                col_widths=[
                    width
                    for _ in po_rows[0]
                ],
            )
        )

    story.append(Spacer(1, 7 * mm))

    # --------------------------------------------------------
    # CO-PSO
    # --------------------------------------------------------

    story.append(
        Table(
            [["CO-PSO Mapping"]],
            colWidths=[260 * mm],
            style=TableStyle([
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    12,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]),
        )
    )

    pso_rows = _table_rows(
        co_pso,
        co_codes,
        pso_codes,
    )

    if len(pso_rows[0]) > 1:
        width = 255 * mm / len(pso_rows[0])
        story.append(
            _pdf_table(
                pso_rows,
                col_widths=[
                    width
                    for _ in pso_rows[0]
                ],
            )
        )

    story.append(Spacer(1, 7 * mm))

    # --------------------------------------------------------
    # MAPPING SCALE
    # --------------------------------------------------------

    story.append(
        Table(
            [["Mapping Scale"]],
            colWidths=[260 * mm],
            style=TableStyle([
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, -1),
                    "Helvetica-Bold",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    12,
                ),
            ]),
        )
    )

    scale_rows = [
        ["Level", "Meaning"],
        ["1", "Low"],
        ["2", "Moderate"],
        ["3", "High"],
        ["Blank", "No meaningful contribution"],
    ]

    story.append(
        _pdf_table(
            scale_rows,
            col_widths=[
                35 * mm,
                90 * mm,
            ],
        )
    )

    doc.build(story)

    return output_path
