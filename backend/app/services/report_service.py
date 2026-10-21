"""
Reference-style OBE attainment report generator.

The downloadable report follows the supplied reference structure.

Student Marks Excel is the authoritative input for:
- assessment/CO mapping
- maximum marks
- student marks

Student-wise marks are used internally by the attainment calculation
and are NOT duplicated into the generated report.

Generated Excel contains exactly two sheets:
1. CO Attainment IA
2. CO Theory Attainment (2)

Important:
The Course Exit Survey produces ONE course-level indirect attainment
value. That same value is displayed consistently for CO, PO and PSO
indirect attainment.
"""


import os
from typing import Any, Dict, List


from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter


from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
    PageBreak,
)


# ============================================================
# STYLES
# ============================================================

THIN = Side(
    style="thin",
    color="000000",
)

BORDER = Border(
    left=THIN,
    right=THIN,
    top=THIN,
    bottom=THIN,
)


# ============================================================
# BASIC HELPERS
# ============================================================

def _f(
    value: Any,
    default: float = 0.0,
) -> float:
    """Safely convert a value to float."""

    try:

        if value is None or value == "":
            return default

        return float(value)

    except Exception:

        return default


def _fmt(value: Any):
    """
    Format values for Excel/PDF.

    Floats are rounded to 2 decimal places for normal
    report display.
    """

    if value is None or value == "":
        return ""

    if isinstance(value, float):
        return round(value, 2)

    return value


def _safe(value: Any) -> str:
    """Safely convert a value to string."""

    return str(
        value
        if value is not None
        else ""
    )


# ============================================================
# EXCEL TABLE STYLE
# ============================================================

def _apply_table_style(
    ws,
    min_row,
    max_row,
    min_col,
    max_col,
    header_rows=None,
    fill="D9EAF7",
):
    """Apply borders, alignment and header formatting."""

    header_rows = header_rows or []

    for row in ws.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):

        for cell in row:

            cell.border = BORDER

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    for row_num in header_rows:

        for col in range(
            min_col,
            max_col + 1,
        ):

            cell = ws.cell(
                row_num,
                col,
            )

            cell.font = Font(
                bold=True
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor=fill,
            )


# ============================================================
# EXCEL COLUMN WIDTHS
# ============================================================

def _set_widths(ws):

    for column in ws.columns:

        letter = get_column_letter(
            column[0].column
        )

        maximum = 0

        for cell in column:

            maximum = max(
                maximum,
                len(
                    str(
                        cell.value or ""
                    )
                ),
            )

        ws.column_dimensions[
            letter
        ].width = min(
            max(
                12,
                maximum + 2,
            ),
            32,
        )


# ============================================================
# ORDERED ASSESSMENT COLUMNS
# ============================================================

def _ordered_assessment_columns(
    result: Dict[str, Any],
) -> List[Dict[str, Any]]:
    """
    Create the assessment/CO columns used in the
    reference-style threshold summary.

    Non-ESE assessments:
        one column per mapped CO.

    ESE:
        represented once because the faculty Excel
        supplies one ESE total mark column mapped
        to CO1-CO6.
    """

    assessments = (
        result.get(
            "assessments",
            [],
        )
        or []
    )

    statistics = (
        result.get(
            "assessment_statistics",
            {},
        )
        or {}
    )

    output: List[
        Dict[str, Any]
    ] = []

    for assessment in assessments:

        name = _safe(
            assessment.get(
                "name",
                "Assessment",
            )
        )

        name_lower = name.lower()

        is_ese = (
            "ese" in name_lower
            or "end sem" in name_lower
            or "endsem" in name_lower
        )

        maximums = (
            assessment.get(
                "co_max_marks",
                {},
            )
            or {}
        )

        mapping = (
            assessment.get(
                "mapping",
                [],
            )
            or []
        )

        # ----------------------------------------------------
        # ESE
        # ----------------------------------------------------

        if is_ese:

            first_co = next(
                iter(
                    maximums.keys()
                ),
                None,
            )

            stat = (
                (
                    statistics.get(
                        name,
                        {},
                    )
                    or {}
                ).get(
                    first_co,
                    {},
                )
                if first_co
                else {}
            )

            max_marks = (
                maximums.get(
                    first_co,
                    0,
                )
                if first_co
                else 0
            )

            output.append(
                {
                    "assessment": name,
                    "co": "CO1-CO6",
                    "max_marks": max_marks,
                    "stat": stat,
                }
            )

            continue

        # ----------------------------------------------------
        # NON-ESE
        # ----------------------------------------------------

        cos = list(
            maximums.keys()
        )

        if not cos:

            for item in mapping:

                for co in (
                    item.get(
                        "cos",
                        [],
                    )
                    or []
                ):

                    if co not in cos:
                        cos.append(co)

        for co in cos:

            output.append(
                {
                    "assessment": name,
                    "co": str(co),
                    "max_marks": maximums.get(
                        co,
                        "",
                    ),
                    "stat": (
                        (
                            statistics.get(
                                name,
                                {},
                            )
                            or {}
                        ).get(
                            co,
                            {},
                        )
                    ),
                }
            )

    return output


# ============================================================
# SINGLE COURSE INDIRECT ATTAINMENT
# ============================================================

def _survey_indirect_value(
    result: Dict[str, Any],
) -> float:
    """
    Return the SINGLE course-level indirect attainment.

    PRIMARY SOURCE:
        result["course_indirect_attainment"]

    This value is calculated by attainment_calculator.py
    from the Course Exit Survey.

    Example:

        Survey = 92.1839%

        Indirect attainment
            = 92.1839 / 100 * 3
            = 2.76552

    This function intentionally does NOT calculate a separate
    indirect value for CO, PO or PSO.

    One course -> one indirect attainment value.
    """

    # --------------------------------------------------------
    # PRIMARY SOURCE
    # --------------------------------------------------------

    course_indirect = result.get(
        "course_indirect_attainment"
    )

    if course_indirect not in (
        None,
        "",
    ):

        return round(
            _f(course_indirect),
            5,
        )

    # --------------------------------------------------------
    # COMPATIBILITY FALLBACK
    #
    # This supports older result structures.
    # --------------------------------------------------------

    indirect = (
        result.get(
            "indirect_co_attainment",
            {},
        )
        or {}
    )

    details = (
        result.get(
            "survey_details",
            {},
        )
        or {}
    )

    overall_pct = details.get(
        "overall_percentage"
    )

    if overall_pct not in (
        None,
        "",
    ):

        return round(
            _f(overall_pct)
            / 100.0
            * 3.0,
            5,
        )

    # --------------------------------------------------------
    # Search existing aggregate value
    # --------------------------------------------------------

    for item in indirect.values():

        if (
            isinstance(
                item,
                dict,
            )
            and item.get(
                "survey_overall_level"
            )
            not in (
                None,
                "",
            )
        ):

            return round(
                _f(
                    item.get(
                        "survey_overall_level"
                    )
                ),
                5,
            )

    # --------------------------------------------------------
    # Last fallback:
    # average CO survey percentages
    # --------------------------------------------------------

    percentages = []

    for item in indirect.values():

        if (
            isinstance(
                item,
                dict,
            )
            and item.get(
                "survey_percentage"
            )
            not in (
                None,
                "",
            )
        ):

            percentages.append(
                _f(
                    item.get(
                        "survey_percentage"
                    )
                )
            )

    if percentages:

        return round(
            (
                sum(percentages)
                / len(percentages)
            )
            / 100.0
            * 3.0,
            5,
        )

    return 0.0


# ============================================================
# ORDERED CO CODES
# ============================================================

def _ordered_co_codes(
    result: Dict[str, Any],
) -> List[str]:
    """
    Keep assessment order first, then append
    remaining CO codes.
    """

    result_codes = [
        str(x)
        for x in (
            result.get(
                "co_codes",
                [],
            )
            or []
        )
    ]

    ordered: List[str] = []

    for item in _ordered_assessment_columns(
        result
    ):

        co = item["co"]

        if co == "CO1-CO6":
            continue

        if co not in ordered:
            ordered.append(co)

    for co in result_codes:

        if co not in ordered:
            ordered.append(co)

    return ordered


# ============================================================
# CO-PO / CO-PSO ARTICULATION ROWS
# ============================================================

def _mapping_rows(
    result: Dict[str, Any],
    co_codes: List[str],
    po_codes: List[str],
    pso_codes: List[str],
):

    co_po = (
        result.get(
            "co_po_articulation_matrix",
            {},
        )
        or {}
    )

    co_pso = (
        result.get(
            "co_pso_articulation_matrix",
            {},
        )
        or {}
    )

    overall = (
        result.get(
            "overall_co_attainment",
            {},
        )
        or {}
    )

    rows = []

    for co in co_codes:

        row = [
            co,
            _fmt(
                (
                    overall.get(
                        co,
                        {},
                    )
                    or {}
                ).get(
                    "overall",
                    "",
                )
            ),
        ]

        for po in po_codes:

            value = (
                (
                    co_po.get(
                        co,
                        {},
                    )
                    or {}
                ).get(
                    po,
                    "",
                )
            )

            row.append(
                _fmt(value)
                if value
                not in (
                    None,
                    0,
                )
                else ""
            )

        for pso in pso_codes:

            value = (
                (
                    co_pso.get(
                        co,
                        {},
                    )
                    or {}
                ).get(
                    pso,
                    "",
                )
            )

            row.append(
                _fmt(value)
                if value
                not in (
                    None,
                    0,
                )
                else ""
            )

        rows.append(row)

    return rows


# ============================================================
# COURSE HEADER
# ============================================================

def _write_course_header(
    ws,
    course: Dict[str, Any],
    last_col: int,
    title: str,
):

    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=last_col,
    )

    ws.cell(
        1,
        1,
        title,
    ).font = Font(
        bold=True,
        size=16,
    )

    ws.cell(
        1,
        1,
    ).alignment = Alignment(
        horizontal="center"
    )

    ws.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=last_col,
    )

    ws.cell(
        2,
        1,
        (
            f"Course: "
            f"{_safe(course.get('course_code'))} - "
            f"{_safe(course.get('course_name'))} | "
            f"Semester: "
            f"{_safe(course.get('semester'))} | "
            f"Academic Year: "
            f"{_safe(course.get('academic_year'))}"
        ),
    ).font = Font(
        bold=True
    )

    ws.cell(
        2,
        1,
    ).alignment = Alignment(
        horizontal="center"
    )


# ============================================================
# GENERATE EXCEL
# ============================================================

def generate_attainment_excel(
    result: Dict[str, Any],
    output_path: str,
):

    wb = Workbook()

    ws = wb.active

    ws.title = "CO Attainment IA"

    # --------------------------------------------------------
    # BASIC DATA
    # --------------------------------------------------------

    course = (
        result.get(
            "course",
            {},
        )
        or {}
    )

    co_codes = _ordered_co_codes(
        result
    )

    po_codes = [
        str(x)
        for x in (
            result.get(
                "po_codes",
                [],
            )
            or []
        )
    ]

    pso_codes = [
        str(x)
        for x in (
            result.get(
                "pso_codes",
                [],
            )
            or []
        )
    ]

    overall = (
        result.get(
            "overall_co_attainment",
            {},
        )
        or {}
    )

    overall_ia = (
        result.get(
            "overall_ia",
            {},
        )
        or {}
    )

    assessment_columns = (
        _ordered_assessment_columns(
            result
        )
    )

    # --------------------------------------------------------
    # SINGLE INDIRECT VALUE
    # --------------------------------------------------------

    indirect_value = (
        _survey_indirect_value(
            result
        )
    )

    # --------------------------------------------------------
    # COURSE HEADER
    # --------------------------------------------------------

    last_col = max(
        8,
        len(
            assessment_columns
        ) + 2,
    )

    _write_course_header(
        ws,
        course,
        last_col,
        "Course Outcome (CO)- Attainment",
    )

    # ========================================================
    # 1. THRESHOLD SUMMARY
    # ========================================================

    row = 4

    ws.cell(
        row,
        1,
        "Assessment",
    )

    ws.cell(
        row,
        2,
        "CO",
    )

    for idx, item in enumerate(
        assessment_columns,
        start=3,
    ):

        ws.cell(
            row,
            idx,
            item["assessment"],
        )

    _apply_table_style(
        ws,
        row,
        row,
        1,
        max(
            2,
            len(
                assessment_columns
            ) + 2,
        ),
        header_rows=[row],
    )

    row += 1

    ws.cell(
        row,
        1,
        "CO",
    )

    ws.cell(
        row,
        2,
        "",
    )

    for idx, item in enumerate(
        assessment_columns,
        start=3,
    ):

        ws.cell(
            row,
            idx,
            item["co"],
        )

    row += 1

    ws.cell(
        row,
        1,
        "Maximum Marks",
    )

    ws.cell(
        row,
        2,
        "",
    )

    for idx, item in enumerate(
        assessment_columns,
        start=3,
    ):

        ws.cell(
            row,
            idx,
            _fmt(
                item["max_marks"]
            ),
        )

    row += 1

    summary_rows = [
        (
            "Threshold Marks",
            "threshold",
        ),
        (
            "Total No. of Students",
            "total_students",
        ),
        (
            "No. of students above threshold",
            "above_threshold",
        ),
        (
            "Percentage",
            "percentage",
        ),
        (
            "Level Attained",
            "level",
        ),
    ]

    for label, key in summary_rows:

        ws.cell(
            row,
            1,
            label,
        )

        ws.cell(
            row,
            2,
            "",
        )

        for idx, item in enumerate(
            assessment_columns,
            start=3,
        ):

            ws.cell(
                row,
                idx,
                _fmt(
                    (
                        item.get(
                            "stat",
                            {},
                        )
                        or {}
                    ).get(
                        key,
                        "",
                    )
                ),
            )

        row += 1

    summary_end = row - 1

    _apply_table_style(
        ws,
        4,
        summary_end,
        1,
        len(
            assessment_columns
        ) + 2,
        header_rows=[
            4,
            5,
            6,
        ],
    )

    # ========================================================
    # 2. INDIRECT ATTAINMENT
    # ========================================================

    row += 1

    ws.cell(
        row,
        1,
        "Indirect Attainment (Course Exit Survey) :",
    )

    ws.cell(
        row,
        2,
        indirect_value,
    )

    ws.cell(
        row,
        1,
    ).font = Font(
        bold=True
    )

    ws.cell(
        row,
        2,
    ).font = Font(
        bold=True
    )

    row += 2

    # ========================================================
    # 3. CO ATTAINMENT
    # ========================================================

    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=5,
    )

    ws.cell(
        row,
        1,
        "CO- Attainment (Internal+External)",
    ).font = Font(
        bold=True,
        size=13,
    )

    ws.cell(
        row,
        1,
    ).alignment = Alignment(
        horizontal="center"
    )

    row += 1

    headers = [
        "CO",
        "Indirect",
        "Direct",
        "Overall IA",
        "O.A. (I+E)",
    ]

    for col, header in enumerate(
        headers,
        start=1,
    ):

        ws.cell(
            row,
            col,
            header,
        )

    header_row = row

    row += 1

    for co in co_codes:

        ia = (
            overall_ia.get(
                co,
                {},
            )
            or {}
        )

        final = (
            overall.get(
                co,
                {},
            )
            or {}
        )

        ws.cell(
            row,
            1,
            co,
        )

        # IMPORTANT:
        # Always use the single survey-derived value.
        ws.cell(
            row,
            2,
            indirect_value,
        )

        ws.cell(
            row,
            3,
            _fmt(
                ia.get(
                    "direct",
                    final.get(
                        "direct",
                        "",
                    ),
                )
            ),
        )

        ws.cell(
            row,
            4,
            _fmt(
                ia.get(
                    "overall_ia",
                    final.get(
                        "overall_ia",
                        "",
                    ),
                )
            ),
        )

        ws.cell(
            row,
            5,
            _fmt(
                final.get(
                    "overall",
                    final.get(
                        "overall_attainment",
                        "",
                    ),
                )
            ),
        )

        row += 1

    _apply_table_style(
        ws,
        header_row,
        row - 1,
        1,
        5,
        header_rows=[
            header_row
        ],
    )

    # ========================================================
    # CO SUMMARY
    # ========================================================

    row += 1

    average_co = (
        sum(
            _f(
                (
                    overall.get(
                        co,
                        {},
                    )
                    or {}
                ).get(
                    "overall",
                    0,
                )
            )
            for co in co_codes
        )
        / len(co_codes)
        if co_codes
        else 0
    )

    target_value = (
        _f(
            (
                overall.get(
                    co_codes[0],
                    {},
                )
                or {}
            ).get(
                "target",
                1.2,
            )
        )
        if co_codes
        else 1.2
    )

    ws.cell(
        row,
        3,
        "Overall CO Attained",
    ).font = Font(
        bold=True
    )

    ws.cell(
        row,
        4,
        _fmt(
            average_co
        ),
    ).font = Font(
        bold=True
    )

    row += 1

    ws.cell(
        row,
        3,
        "CO Target",
    ).font = Font(
        bold=True
    )

    ws.cell(
        row,
        4,
        _fmt(
            target_value
        ),
    ).font = Font(
        bold=True
    )

    row += 1

    ws.cell(
        row,
        3,
        "CO Attain or Not",
    ).font = Font(
        bold=True
    )

    ws.cell(
        row,
        4,
        (
            "Yes"
            if average_co
            >= target_value
            else "No"
        ),
    ).font = Font(
        bold=True
    )

    _apply_table_style(
        ws,
        row - 2,
        row,
        3,
        4,
    )

    # --------------------------------------------------------
    # No extra methodology/chart/student table
    # --------------------------------------------------------

    _set_widths(ws)

    ws.freeze_panes = "A7"

    # ========================================================
    # SECOND SHEET
    # ========================================================

    ws2 = wb.create_sheet(
        "CO Theory Attainment (2)"
    )

    total_cols = max(
        3,
        2
        + len(po_codes)
        + len(pso_codes),
    )

    _write_course_header(
        ws2,
        course,
        total_cols,
        "CO Theory Attainment (2)",
    )

    # ========================================================
    # CO-PO/PSO ARTICULATION MATRIX
    # ========================================================

    row = 4

    ws2.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=total_cols,
    )

    ws2.cell(
        row,
        1,
        "CO-PO Articulation Matrix",
    ).font = Font(
        bold=True,
        size=13,
    )

    ws2.cell(
        row,
        1,
    ).alignment = Alignment(
        horizontal="center"
    )

    row += 1

    headers = (
        ["CO", "CO Attainment"]
        + po_codes
        + pso_codes
    )

    for col, value in enumerate(
        headers,
        start=1,
    ):

        ws2.cell(
            row,
            col,
            value,
        )

    matrix_header = row

    row += 1

    for values in _mapping_rows(
        result,
        co_codes,
        po_codes,
        pso_codes,
    ):

        for col, value in enumerate(
            values,
            start=1,
        ):

            ws2.cell(
                row,
                col,
                value,
            )

        row += 1

    # ========================================================
    # AVERAGE ARTICULATION
    # ========================================================

    ws2.cell(
        row,
        1,
        "Avg CO",
    )

    avg_co = (
        sum(
            _f(
                (
                    overall.get(
                        co,
                        {},
                    )
                    or {}
                ).get(
                    "overall",
                    0,
                )
            )
            for co in co_codes
        )
        / len(co_codes)
        if co_codes
        else 0
    )

    ws2.cell(
        row,
        2,
        _fmt(
            avg_co
        ),
    )

    for idx, outcome in enumerate(
        po_codes + pso_codes,
        start=3,
    ):

        if outcome in po_codes:

            mapping = (
                result.get(
                    "co_po_articulation_matrix",
                    {},
                )
                or {}
            )

        else:

            mapping = (
                result.get(
                    "co_pso_articulation_matrix",
                    {},
                )
                or {}
            )

        vals = []

        for co in co_codes:

            value = (
                (
                    mapping.get(
                        co,
                        {},
                    )
                    or {}
                ).get(
                    outcome
                )
            )

            if value not in (
                None,
                "",
                0,
            ):

                vals.append(
                    _f(value)
                )

        ws2.cell(
            row,
            idx,
            _fmt(
                (
                    sum(vals)
                    / len(vals)
                    if vals
                    else ""
                )
            ),
        )

    row += 2

    _apply_table_style(
        ws2,
        matrix_header,
        row - 3,
        1,
        total_cols,
        header_rows=[
            matrix_header
        ],
    )

    # ========================================================
    # PO ATTAINMENT
    # ========================================================

    ws2.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=max(
            2,
            len(po_codes) + 1,
        ),
    )

    ws2.cell(
        row,
        1,
        "CO-PO Attainment Matrix",
    ).font = Font(
        bold=True,
        size=13,
    )

    ws2.cell(
        row,
        1,
    ).alignment = Alignment(
        horizontal="center"
    )

    row += 1

    for col, value in enumerate(
        [""] + po_codes,
        start=1,
    ):

        ws2.cell(
            row,
            col,
            value,
        )

    po_header = row

    row += 1

    for label, key in [
        (
            "PO Attained (Direct)",
            "direct_attainment",
        ),
        (
            "PO Attained (Indirect)",
            "indirect_attainment",
        ),
        (
            "Attained",
            "overall_attainment",
        ),
    ]:

        ws2.cell(
            row,
            1,
            label,
        )

        for idx, po in enumerate(
            po_codes,
            start=2,
        ):

            item = (
                (
                    result.get(
                        "po_attainment",
                        {},
                    )
                    or {}
                ).get(
                    po,
                    {},
                )
                or {}
            )

            # IMPORTANT:
            # PO indirect always equals the
            # Course Exit Survey indirect value.
            if key == "indirect_attainment":

                value = indirect_value

            else:

                value = item.get(
                    key,
                    "",
                )

            ws2.cell(
                row,
                idx,
                _fmt(value),
            )

        row += 1

    _apply_table_style(
        ws2,
        po_header,
        row - 1,
        1,
        max(
            2,
            len(po_codes) + 1,
        ),
        header_rows=[
            po_header
        ],
    )

    # ========================================================
    # PSO ATTAINMENT
    # ========================================================

    row += 1

    ws2.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=max(
            2,
            len(pso_codes) + 1,
        ),
    )

    ws2.cell(
        row,
        1,
        "CO-PSO Attainment Matrix",
    ).font = Font(
        bold=True,
        size=13,
    )

    ws2.cell(
        row,
        1,
    ).alignment = Alignment(
        horizontal="center"
    )

    row += 1

    for col, value in enumerate(
        [""] + pso_codes,
        start=1,
    ):

        ws2.cell(
            row,
            col,
            value,
        )

    pso_header = row

    row += 1

    for label, key in [
        (
            "PSO Attained (Direct)",
            "direct_attainment",
        ),
        (
            "PSO Attained (Indirect)",
            "indirect_attainment",
        ),
        (
            "Attained",
            "overall_attainment",
        ),
    ]:

        ws2.cell(
            row,
            1,
            label,
        )

        for idx, pso in enumerate(
            pso_codes,
            start=2,
        ):

            item = (
                (
                    result.get(
                        "pso_attainment",
                        {},
                    )
                    or {}
                ).get(
                    pso,
                    {},
                )
                or {}
            )

            # IMPORTANT:
            # PSO indirect always equals the
            # Course Exit Survey indirect value.
            if key == "indirect_attainment":

                value = indirect_value

            else:

                value = item.get(
                    key,
                    "",
                )

            ws2.cell(
                row,
                idx,
                _fmt(value),
            )

        row += 1

    if pso_codes:

        _apply_table_style(
            ws2,
            pso_header,
            row - 1,
            1,
            len(
                pso_codes
            ) + 1,
            header_rows=[
                pso_header
            ],
        )

    _set_widths(ws2)

    ws2.freeze_panes = "A6"

    # ========================================================
    # SAVE EXCEL
    # ========================================================

    os.makedirs(
        os.path.dirname(
            output_path
        )
        or ".",
        exist_ok=True,
    )

    wb.save(
        output_path
    )

    return output_path


# ============================================================
# COMPATIBILITY WRAPPER
# ============================================================

def generate_excel(
    result,
    output_path,
):

    return generate_attainment_excel(
        result,
        output_path,
    )


# ============================================================
# PDF HELPERS
# ============================================================

def _pdf_text(
    value: Any,
) -> str:

    return (
        _safe(value)
        .replace(
            "&",
            "&amp;",
        )
        .replace(
            "<",
            "&lt;",
        )
        .replace(
            ">",
            "&gt;",
        )
    )


def _pdf_table(
    rows,
    font_size=7,
):

    table = Table(
        rows,
        repeatRows=1,
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.4,
                    colors.black,
                ),
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.lightgrey,
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold",
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, -1),
                    "CENTER",
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    font_size,
                ),
            ]
        )
    )

    return table


# ============================================================
# GENERATE PDF
# ============================================================

def generate_attainment_pdf(
    result: Dict[str, Any],
    output_path: str,
):

    course = (
        result.get(
            "course",
            {},
        )
        or {}
    )

    co_codes = _ordered_co_codes(
        result
    )

    po_codes = [
        str(x)
        for x in (
            result.get(
                "po_codes",
                [],
            )
            or []
        )
    ]

    pso_codes = [
        str(x)
        for x in (
            result.get(
                "pso_codes",
                [],
            )
            or []
        )
    ]

    overall = (
        result.get(
            "overall_co_attainment",
            {},
        )
        or {}
    )

    overall_ia = (
        result.get(
            "overall_ia",
            {},
        )
        or {}
    )

    assessment_columns = (
        _ordered_assessment_columns(
            result
        )
    )

    # --------------------------------------------------------
    # SINGLE INDIRECT VALUE
    # --------------------------------------------------------

    indirect_value = (
        _survey_indirect_value(
            result
        )
    )

    # --------------------------------------------------------
    # PDF DOCUMENT
    # --------------------------------------------------------

    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(A4),
        rightMargin=8 * mm,
        leftMargin=8 * mm,
        topMargin=8 * mm,
        bottomMargin=8 * mm,
    )

    styles = (
        getSampleStyleSheet()
    )

    elements = []

    # ========================================================
    # TITLE
    # ========================================================

    elements.append(
        Paragraph(
            "<b>FACULTY AGENT</b>",
            styles["Title"],
        )
    )

    elements.append(
        Paragraph(
            "Course Outcome (CO)- Attainment",
            styles["Heading2"],
        )
    )

    elements.append(
        Paragraph(
            (
                f"Course: "
                f"{_pdf_text(course.get('course_code'))} - "
                f"{_pdf_text(course.get('course_name'))} | "
                f"Semester: "
                f"{_pdf_text(course.get('semester'))} | "
                f"Academic Year: "
                f"{_pdf_text(course.get('academic_year'))}"
            ),
            styles["Normal"],
        )
    )

    elements.append(
        Spacer(
            1,
            5,
        )
    )

    # ========================================================
    # THRESHOLD SUMMARY
    # ========================================================

    threshold_rows = [
        ["Metric"]
        + [
            f"{x['assessment']}\n{x['co']}"
            for x in assessment_columns
        ]
    ]

    for label, key in [
        (
            "Threshold Marks",
            "threshold",
        ),
        (
            "Total No. of Students",
            "total_students",
        ),
        (
            "No. of students above threshold",
            "above_threshold",
        ),
        (
            "Percentage",
            "percentage",
        ),
        (
            "Level Attained",
            "level",
        ),
    ]:

        threshold_rows.append(
            [label]
            + [
                _fmt(
                    (
                        x.get(
                            "stat",
                            {},
                        )
                        or {}
                    ).get(
                        key,
                        "",
                    )
                )
                for x in assessment_columns
            ]
        )

    elements.append(
        _pdf_table(
            threshold_rows,
            font_size=6,
        )
    )

    elements.append(
        Spacer(
            1,
            5,
        )
    )

    # ========================================================
    # INDIRECT ATTAINMENT
    # ========================================================

    elements.append(
        Paragraph(
            (
                "<b>Indirect Attainment "
                "(Course Exit Survey): "
                f"{_fmt(indirect_value)}</b>"
            ),
            styles["Normal"],
        )
    )

    elements.append(
        Spacer(
            1,
            7,
        )
    )

    # ========================================================
    # CO ATTAINMENT
    # ========================================================

    co_rows = [
        [
            "CO",
            "Indirect",
            "Direct",
            "Overall IA",
            "O.A. (I+E)",
        ]
    ]

    for co in co_codes:

        ia = (
            overall_ia.get(
                co,
                {},
            )
            or {}
        )

        final = (
            overall.get(
                co,
                {},
            )
            or {}
        )

        co_rows.append(
            [
                co,

                # IMPORTANT:
                # Single survey-derived value.
                _fmt(
                    indirect_value
                ),

                _fmt(
                    ia.get(
                        "direct",
                        final.get(
                            "direct",
                            "",
                        ),
                    )
                ),

                _fmt(
                    ia.get(
                        "overall_ia",
                        final.get(
                            "overall_ia",
                            "",
                        ),
                    )
                ),

                _fmt(
                    final.get(
                        "overall",
                        final.get(
                            "overall_attainment",
                            "",
                        ),
                    )
                ),
            ]
        )

    elements.append(
        _pdf_table(
            co_rows,
            font_size=7,
        )
    )

    # ========================================================
    # CO SUMMARY
    # ========================================================

    avg = (
        sum(
            _f(
                (
                    overall.get(
                        co,
                        {},
                    )
                    or {}
                ).get(
                    "overall",
                    0,
                )
            )
            for co in co_codes
        )
        / len(co_codes)
        if co_codes
        else 0
    )

    target = (
        _f(
            (
                overall.get(
                    co_codes[0],
                    {},
                )
                or {}
            ).get(
                "target",
                1.2,
            )
        )
        if co_codes
        else 1.2
    )

    elements.append(
        Spacer(
            1,
            5,
        )
    )

    elements.append(
        _pdf_table(
            [
                [
                    "Overall CO Attained",
                    "CO Target",
                    "CO Attain or Not",
                ],
                [
                    _fmt(avg),
                    _fmt(target),
                    (
                        "Yes"
                        if avg >= target
                        else "No"
                    ),
                ],
            ],
            font_size=7,
        )
    )

    # ========================================================
    # SECOND PAGE
    # ========================================================

    elements.append(
        PageBreak()
    )

    elements.append(
        Paragraph(
            "CO Theory Attainment (2)",
            styles["Heading2"],
        )
    )

    # ========================================================
    # CO-PO/PSO ARTICULATION MATRIX
    # ========================================================

    co_po = (
        result.get(
            "co_po_articulation_matrix",
            {},
        )
        or {}
    )

    co_pso = (
        result.get(
            "co_pso_articulation_matrix",
            {},
        )
        or {}
    )

    matrix_rows = [
        [
            "CO",
            "CO Attainment",
        ]
        + po_codes
        + pso_codes
    ]

    for co in co_codes:

        row = [
            co,
            _fmt(
                (
                    overall.get(
                        co,
                        {},
                    )
                    or {}
                ).get(
                    "overall",
                    "",
                )
            ),
        ]

        row += [
            _fmt(
                (
                    co_po.get(
                        co,
                        {},
                    )
                    or {}
                ).get(
                    po,
                    "",
                )
            )
            for po in po_codes
        ]

        row += [
            _fmt(
                (
                    co_pso.get(
                        co,
                        {},
                    )
                    or {}
                ).get(
                    pso,
                    "",
                )
            )
            for pso in pso_codes
        ]

        matrix_rows.append(
            row
        )

    elements.append(
        Paragraph(
            "CO-PO Articulation Matrix",
            styles["Heading3"],
        )
    )

    elements.append(
        _pdf_table(
            matrix_rows,
            font_size=5,
        )
    )

    elements.append(
        Spacer(
            1,
            8,
        )
    )

    # ========================================================
    # PO ATTAINMENT
    # ========================================================

    po_rows = [
        [
            "PO",
            "Direct",
            "Indirect",
            "Attained",
        ]
    ]

    for po in po_codes:

        item = (
            (
                result.get(
                    "po_attainment",
                    {},
                )
                or {}
            ).get(
                po,
                {},
            )
            or {}
        )

        po_rows.append(
            [
                po,

                _fmt(
                    item.get(
                        "direct_attainment",
                        item.get(
                            "attainment",
                            "",
                        ),
                    )
                ),

                # IMPORTANT:
                # Always use the same
                # Course Exit Survey value.
                _fmt(
                    indirect_value
                ),

                _fmt(
                    item.get(
                        "overall_attainment",
                        "",
                    )
                ),
            ]
        )

    if len(po_rows) > 1:

        elements.append(
            Paragraph(
                "CO-PO Attainment Matrix",
                styles["Heading3"],
            )
        )

        elements.append(
            _pdf_table(
                po_rows,
                font_size=6,
            )
        )

        elements.append(
            Spacer(
                1,
                8,
            )
        )

    # ========================================================
    # PSO ATTAINMENT
    # ========================================================

    pso_rows = [
        [
            "PSO",
            "Direct",
            "Indirect",
            "Attained",
        ]
    ]

    for pso in pso_codes:

        item = (
            (
                result.get(
                    "pso_attainment",
                    {},
                )
                or {}
            ).get(
                pso,
                {},
            )
            or {}
        )

        pso_rows.append(
            [
                pso,

                _fmt(
                    item.get(
                        "direct_attainment",
                        item.get(
                            "attainment",
                            "",
                        ),
                    )
                ),

                # IMPORTANT:
                # Always use the same
                # Course Exit Survey value.
                _fmt(
                    indirect_value
                ),

                _fmt(
                    item.get(
                        "overall_attainment",
                        "",
                    )
                ),
            ]
        )

    if len(pso_rows) > 1:

        elements.append(
            Paragraph(
                "CO-PSO Attainment Matrix",
                styles["Heading3"],
            )
        )

        elements.append(
            _pdf_table(
                pso_rows,
                font_size=6,
            )
        )

    # ========================================================
    # SAVE PDF
    # ========================================================

    os.makedirs(
        os.path.dirname(
            output_path
        )
        or ".",
        exist_ok=True,
    )

    doc.build(
        elements
    )

    return output_path


# ============================================================
# COMPATIBILITY WRAPPER
# ============================================================

def generate_pdf(
    result,
    output_path,
):

    return generate_attainment_pdf(
        result,
        output_path,
    )