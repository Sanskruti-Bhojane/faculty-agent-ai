import os
import pandas as pd

from pypdf import PdfReader
from docx import Document


def extract_pdf(file_path: str) -> str:
    """Extract text from PDF."""

    reader = PdfReader(file_path)

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n".join(pages)


def extract_docx(file_path: str) -> str:
    """Extract text from Word document."""

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Also extract tables
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:
                row_text.append(cell.text.strip())

            paragraphs.append(" | ".join(row_text))

    return "\n".join(paragraphs)


def extract_excel(file_path: str) -> str:
    """Extract text/data from Excel."""

    sheets = pd.read_excel(
        file_path,
        sheet_name=None,
        header=None
    )

    output = []

    for sheet_name, dataframe in sheets.items():

        output.append(
            f"\n--- SHEET: {sheet_name} ---\n"
        )

        dataframe = dataframe.fillna("")

        for row in dataframe.values:

            row_text = [
                str(value).strip()
                for value in row
                if str(value).strip()
            ]

            if row_text:
                output.append(
                    " | ".join(row_text)
                )

    return "\n".join(output)


def extract_text(file_path: str) -> str:
    """
    Automatically extract text based on file type.
    """

    if not os.path.exists(file_path):

        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = os.path.splitext(
        file_path
    )[1].lower()


    # PDF
    if extension == ".pdf":

        return extract_pdf(file_path)


    # Word
    elif extension == ".docx":

        return extract_docx(file_path)


    # Excel
    elif extension in [".xlsx", ".xlsm"]:

        return extract_excel(file_path)


    # CSV
    elif extension == ".csv":

        dataframe = pd.read_csv(
            file_path,
            header=None
        )

        dataframe = dataframe.fillna("")

        output = []

        for row in dataframe.values:

            row_text = [
                str(value).strip()
                for value in row
                if str(value).strip()
            ]

            if row_text:
                output.append(
                    " | ".join(row_text)
                )

        return "\n".join(output)


    # TXT
    elif extension == ".txt":

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            return file.read()


    # Images
    elif extension in [
        ".png",
        ".jpg",
        ".jpeg"
    ]:

        # Image extraction will be handled
        # by the multimodal AI pipeline later.
        return (
            "[IMAGE FILE]\n"
            "Image text extraction will be "
            "handled by the AI document agent."
        )


    else:

        raise ValueError(
            f"Unsupported file type: {extension}"
        )